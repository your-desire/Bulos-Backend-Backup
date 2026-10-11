"""
Transcription Accuracy Validation Script
=========================================
Randomly samples 20% of the audio files in scripts/audios/ (across all 5
sub-folders), sends each file to the local STT /transcribe endpoint, and
writes a new spreadsheet with the results.

Audio location : scripts/audios/<subfolder>/<filename>.mp3
Expected output: the filename stem (e.g. "abi.mp3" → expected "abi")
Language       : all audio is Bulos ("bul")
API endpoint   : http://127.0.0.1:8000/api/v1/stt/transcribe

Output spreadsheet columns:
  1. Audio File          — relative path from audios/ (subfolder/filename.mp3)
  2. Subfolder           — which collection the file came from
  3. Expected Output     — filename stem (the actual spoken text)
  4. STT Transcription   — what the model returned
  5. chrF Score          — character n-gram F-score (0–1, see below)

chrF Score formula:
  Harmonic mean of character-level precision and recall over n-gram orders 1–6.
  Score range: 0.0 (no match at all) → 1.0 (exact match)
  Colour coding:
    Green  (≥ 0.70) — good transcription
    Yellow (≥ 0.40) — partial / acceptable
    Red    (< 0.40) — poor
    Grey             — API error

Usage (run locally while the backend is running on 127.0.0.1:8000):
    python scripts/validate_transcription.py

No other project files are touched.
"""

import math
import random
import re
import sys
import time
from collections import Counter
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

SCRIPT_DIR    = Path(__file__).parent
AUDIOS_DIR    = SCRIPT_DIR / "audios"
API_URL       = "http://127.0.0.1:8000/api/v1/stt/transcribe"
DEVICE_ID     = "123e4567-e89b-42d3-a456-426614174000"
LANGUAGE      = "bul"
SAMPLE_RATIO  = 0.20
RANDOM_SEED   = 42
REQUEST_DELAY   = 0.3    # seconds between calls
REQUEST_TIMEOUT = 120    # seconds per call

AUDIO_EXTENSIONS = {".mp3", ".wav", ".ogg", ".flac", ".m4a", ".webm"}

HEADERS = {
    "accept":      "application/json",
    "X-Device-ID": DEVICE_ID,
}

# ---------------------------------------------------------------------------
# Spreadsheet styling
# ---------------------------------------------------------------------------

HEADER_FILL  = PatternFill("solid", fgColor="2F5496")   # dark blue
HEADER_FONT  = Font(bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

FILL_GOOD    = PatternFill("solid", fgColor="C6EFCE")   # green  ≥ 0.70
FILL_OK      = PatternFill("solid", fgColor="FFEB9C")   # yellow ≥ 0.40
FILL_BAD     = PatternFill("solid", fgColor="FFC7CE")   # red    < 0.40
FILL_ERROR   = PatternFill("solid", fgColor="D9D9D9")   # grey  (API error)

DATA_ALIGN   = Alignment(wrap_text=True, vertical="top")

# ---------------------------------------------------------------------------
# chrF metric  (same implementation as validate_translation.py)
# ---------------------------------------------------------------------------

def _char_ngrams(text: str, n: int) -> Counter:
    text = text.strip()
    return Counter(text[i:i + n] for i in range(max(0, len(text) - n + 1)))


def chrf_score(hypothesis: str, reference: str, n: int = 6, beta: float = 1.0) -> float:
    """
    Character n-gram F-score (Popović, 2015).

    chrP = (matched char-ngrams in hyp) / (total char-ngrams in hyp)
    chrR = (matched char-ngrams in hyp) / (total char-ngrams in ref)
    chrF = (1 + β²) × chrP × chrR / (β² × chrP + chrR)

    β=1 weights precision and recall equally.
    Returns a float in [0.0, 1.0].
    """
    hyp = hypothesis.lower().strip()
    ref = reference.lower().strip()

    if not hyp or not ref:
        return 0.0
    if hyp == ref:
        return 1.0

    prec_total = prec_match = rec_total = rec_match = 0
    for order in range(1, n + 1):
        h_ng = _char_ngrams(hyp, order)
        r_ng = _char_ngrams(ref, order)
        overlap = sum(min(h_ng[k], r_ng[k]) for k in h_ng)
        prec_total += sum(h_ng.values())
        prec_match += overlap
        rec_total  += sum(r_ng.values())
        rec_match  += overlap

    chrP = prec_match / prec_total if prec_total else 0.0
    chrR = rec_match  / rec_total  if rec_total  else 0.0
    if chrP + chrR == 0:
        return 0.0
    b2 = beta ** 2
    return round((1 + b2) * chrP * chrR / (b2 * chrP + chrR), 4)

# ---------------------------------------------------------------------------
# API call
# ---------------------------------------------------------------------------

def transcribe(audio_path: Path, retries: int = 3) -> str:
    """
    POST the audio file to the STT endpoint.
    Returns the transcribed text, or 'ERROR: ...' on failure.
    """
    for attempt in range(1, retries + 1):
        try:
            with open(audio_path, "rb") as f:
                files   = {"audio": (audio_path.name, f, "audio/mpeg")}
                data    = {"language": LANGUAGE}
                resp = requests.post(
                    API_URL,
                    headers=HEADERS,
                    files=files,
                    data=data,
                    timeout=REQUEST_TIMEOUT,
                )
            if resp.status_code == 200:
                return resp.json().get("text", "").strip()
            err = f"ERROR: HTTP {resp.status_code} — {resp.text[:120]}"
            if attempt < retries:
                time.sleep(2 * attempt)
                continue
            return err
        except requests.exceptions.Timeout:
            if attempt < retries:
                time.sleep(5)
                continue
            return "ERROR: timeout"
        except Exception as e:
            return f"ERROR: {e}"
    return "ERROR: all retries failed"

# ---------------------------------------------------------------------------
# Collect & sample audio files
# ---------------------------------------------------------------------------

def collect_audio_files(audios_dir: Path) -> list[Path]:
    """Return all audio files found under audios_dir (any depth)."""
    files = [
        p for p in audios_dir.rglob("*")
        if p.is_file() and p.suffix.lower() in AUDIO_EXTENSIONS
    ]
    return sorted(files)


def sample_files(files: list[Path], ratio: float, seed: int) -> list[Path]:
    rng = random.Random(seed)
    k   = max(1, math.ceil(len(files) * ratio))
    return rng.sample(files, min(k, len(files)))

# ---------------------------------------------------------------------------
# Write spreadsheet
# ---------------------------------------------------------------------------

COL_HEADERS = [
    "Audio File",
    "Subfolder",
    "Expected Output",
    "STT Transcription",
    "chrF Score",
]
COL_WIDTHS = [50, 35, 30, 30, 13]


def write_spreadsheet(results: list[dict], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "STT Validation"

    # Header row
    for ci, (header, width) in enumerate(zip(COL_HEADERS, COL_WIDTHS), 1):
        cell = ws.cell(row=1, column=ci, value=header)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        ws.column_dimensions[get_column_letter(ci)].width = width

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COL_HEADERS))}1"

    # Data rows
    for ri, row in enumerate(results, 2):
        ws.cell(ri, 1, row["audio_file"]).alignment = DATA_ALIGN
        ws.cell(ri, 2, row["subfolder"]).alignment  = DATA_ALIGN
        ws.cell(ri, 3, row["expected"]).alignment   = DATA_ALIGN
        ws.cell(ri, 4, row["transcription"]).alignment = DATA_ALIGN

        chrf_cell = ws.cell(ri, 5)
        chrf_cell.alignment = Alignment(horizontal="center", vertical="top")

        if isinstance(row["chrf"], float):
            chrf_cell.value         = row["chrf"]
            chrf_cell.number_format = "0.000"
            if row["chrf"] >= 0.70:
                chrf_cell.fill = FILL_GOOD
            elif row["chrf"] >= 0.40:
                chrf_cell.fill = FILL_OK
            else:
                chrf_cell.fill = FILL_BAD
        else:
            chrf_cell.value = row["chrf"]   # "API Error" string
            chrf_cell.fill  = FILL_ERROR

    # Aggregate stats below the data
    last_data_row = len(results) + 1
    gap_row       = last_data_row + 2

    valid  = [r for r in results if isinstance(r["chrf"], float)]
    errors = len(results) - len(valid)
    avg    = round(sum(r["chrf"] for r in valid) / len(valid), 4) if valid else "N/A"
    exact  = sum(1 for r in valid
                 if r["transcription"].lower().strip() == r["expected"].lower().strip())
    good   = sum(1 for r in valid if r["chrf"] >= 0.70)
    ok     = sum(1 for r in valid if 0.40 <= r["chrf"] < 0.70)
    poor   = sum(1 for r in valid if r["chrf"] < 0.40)

    stats = [
        ("Total audio files tested", len(results)),
        ("Exact matches",             exact),
        ("Exact match rate",          f"{exact / len(results):.1%}" if results else "N/A"),
        ("Average chrF score",        avg),
        ("chrF ≥ 0.70 (good)",        good),
        ("chrF 0.40–0.69 (ok)",       ok),
        ("chrF < 0.40 (poor)",        poor),
        ("API errors",                errors),
    ]

    stat_header = ws.cell(gap_row, 1, "Aggregate Statistics")
    stat_header.font = Font(bold=True)

    for i, (label, value) in enumerate(stats, gap_row + 1):
        ws.cell(i, 1, label)
        ws.cell(i, 2, value)

    # Per-subfolder breakdown
    by_folder: dict[str, list[float]] = {}
    for r in valid:
        by_folder.setdefault(r["subfolder"], []).append(r["chrf"])

    folder_start = gap_row + len(stats) + 2
    ws.cell(folder_start, 1, "Per-Subfolder Breakdown").font = Font(bold=True)
    ws.cell(folder_start, 2, "Count").font = Font(bold=True)
    ws.cell(folder_start, 3, "Avg chrF").font = Font(bold=True)

    for fi, (folder, scores) in enumerate(sorted(by_folder.items()), folder_start + 1):
        ws.cell(fi, 1, folder)
        ws.cell(fi, 2, len(scores))
        ws.cell(fi, 3, round(sum(scores) / len(scores), 4))

    wb.save(str(out_path))
    print(f"\nSpreadsheet saved to: {out_path}")

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    if not AUDIOS_DIR.exists():
        print(f"ERROR: audios directory not found: {AUDIOS_DIR}")
        sys.exit(1)

    # Collect all audio files
    all_files = collect_audio_files(AUDIOS_DIR)
    if not all_files:
        print(f"ERROR: no audio files found under {AUDIOS_DIR}")
        sys.exit(1)

    print(f"Found {len(all_files)} audio files across "
          f"{len(set(f.parent for f in all_files))} subfolders")

    sample = sample_files(all_files, SAMPLE_RATIO, RANDOM_SEED)
    print(f"Sampled 20%: {len(sample)} files")
    print(f"Calling STT API at: {API_URL}\n")

    results: list[dict] = []
    errors = 0

    for i, audio_path in enumerate(sample, 1):
        expected    = re.sub(r"[^\w\s]", "", audio_path.stem)  # strip punctuation
        subfolder   = audio_path.parent.name   # immediate parent folder name
        relative    = audio_path.relative_to(AUDIOS_DIR)

        sys.stdout.write(
            f"\r[{i}/{len(sample)}]  {subfolder}  '{expected[:40]}'" + " " * 10
        )
        sys.stdout.flush()

        transcription = transcribe(audio_path)
        time.sleep(REQUEST_DELAY)

        is_error = transcription.startswith("ERROR:")
        if is_error:
            errors += 1
            score = "API Error"
        else:
            score = chrf_score(transcription, expected)

        results.append({
            "audio_file":    str(relative),
            "subfolder":     subfolder,
            "expected":      expected,
            "transcription": transcription,
            "chrf":          score,
        })

    print(f"\n\nDone. {errors} API errors out of {len(sample)} files.")

    # Summary
    valid   = [r for r in results if isinstance(r["chrf"], float)]
    avg     = round(sum(r["chrf"] for r in valid) / len(valid), 4) if valid else 0
    exact   = sum(1 for r in valid
                  if r["transcription"].lower().strip() == r["expected"].lower().strip())
    print(f"Exact match : {exact}/{len(results)} ({exact/len(results):.1%})")
    print(f"Avg chrF    : {avg}")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path  = SCRIPT_DIR / f"transcription_validation_{timestamp}.xlsx"
    write_spreadsheet(results, out_path)


if __name__ == "__main__":
    main()
